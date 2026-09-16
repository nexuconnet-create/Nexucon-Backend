"""
Excel bulk upload of PUNDIT readings (review meeting A2).

The operator fills the platform's .xlsx template — one row per test point,
columns in the template's order — and uploads it against a project. Rows are
grouped into tests: an element's CONSECUTIVE rows sharing the same test type
and floor form ONE test with points A, B, C... Every group is created through
``PUNDITTestSerializer`` — the same authoritative write path as the manual
entry form — so velocity, quality grade, E.C.S and crack depth stay
server-computed and non-writable, labels are auto-assigned, and the
element-level means are persisted exactly as they are from the form.

Nothing is committed unless every row validates (all-or-nothing): a statutory
registry is never partially populated from a batch file.
"""
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

# The template's exact column headers. Columns are matched BY HEADER NAME, so
# their order in a filled sheet does not matter — but these are the only
# headers the parser reads.
TEMPLATE_COLUMNS = [
    'STRUCTURAL ELEMENT',
    'FLOOR',
    'TEST TYPE',
    'POINT',
    'PATH LENGTH L (MM)',
    'TRANSIT TIME T (US)',
    'T UNCRACKED (US)',
    'SURFACE CONDITION',
    'TRANSDUCER FREQUENCY (KHZ)',
    'TRANSDUCER TYPE',
    'TEST LOCATION',
    'WEATHER CONDITION',
    'NOTES',
]
REQUIRED_COLUMNS = [
    'STRUCTURAL ELEMENT', 'TEST TYPE', 'PATH LENGTH L (MM)',
    'TRANSIT TIME T (US)', 'T UNCRACKED (US)', 'SURFACE CONDITION',
]
COLUMN_WIDTHS = {
    'STRUCTURAL ELEMENT': 34, 'FLOOR': 16, 'TEST TYPE': 20, 'POINT': 8,
    'PATH LENGTH L (MM)': 14, 'TRANSIT TIME T (US)': 14, 'T UNCRACKED (US)': 14,
    'SURFACE CONDITION': 30, 'TRANSDUCER FREQUENCY (KHZ)': 14,
    'TRANSDUCER TYPE': 16, 'TEST LOCATION': 24, 'WEATHER CONDITION': 20,
    'NOTES': 30,
}

TEST_TYPE_ALIASES = (
    ('pulse', 'pulse_velocity'),
    ('crack', 'crack_depth'),
    ('surface', 'surface_quality'),
    ('homogeneity', 'surface_quality'),
)


def _text(value):
    """Cell text, stripped — None for blank cells."""
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _number(value, row, header, errors):
    """Cell as a float, or None when blank. Non-numeric content is a row
    error — a measurement cell cannot be coerced silently."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool):
        errors.append({'row': row, 'message': f'{header} must be a number.'})
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).strip())
    except ValueError:
        errors.append({'row': row,
                       'message': f'{header} must be a number — got {value!r}.'})
        return None


def _test_type(value, row, errors):
    text = _text(value)
    if text is None:
        errors.append({'row': row,
                       'message': 'TEST TYPE is required (PULSE VELOCITY / CRACK DEPTH / SURFACE QUALITY).'})
        return None
    lowered = text.lower().replace('_', ' ')
    for needle, test_type in TEST_TYPE_ALIASES:
        if needle in lowered:
            return test_type
    errors.append({'row': row,
                   'message': f'TEST TYPE {text!r} is not recognised — use PULSE VELOCITY, '
                              'CRACK DEPTH or SURFACE QUALITY.'})
    return None


def _transducer_type(value):
    text = (_text(value) or 'DIRECT').upper().replace('-', '_').replace(' ', '_')
    if 'SEMI' in text:
        return 'SEMI_DIRECT'
    if 'INDIRECT' in text:
        return 'INDIRECT'
    return 'DIRECT'


def parse_readings_workbook(file_obj):
    """Parse an uploaded template workbook into test groups.

    Returns ``(groups, errors)``. Each group is
    ``{'payload': <PUNDITTestSerializer data>, 'first_row': int,
    'last_row': int, 'element': str, 'test_type': str}``. ``payload['project']``
    is NOT set here — the view injects the scoped project. Errors are
    ``{'row': int, 'message': str}`` (parsing) or ``{'rows': 'a-b', 'message':
    str}`` (grouping); any error means the caller must reject the file.
    """
    try:
        workbook = load_workbook(file_obj, read_only=True, data_only=True)
    except Exception:
        return [], [{'row': 0, 'message': 'The file is not a readable .xlsx workbook — '
                                          'download the template from this page and fill it in.'}]
    try:
        if 'READINGS' not in workbook.sheetnames:
            return [], [{'row': 0,
                         'message': f'The workbook has no "READINGS" sheet (found: '
                                    f'{", ".join(workbook.sheetnames) or "none"}) — '
                                    'download the template from this page.'}]
        sheet = workbook['READINGS']
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            headers = next(rows_iter)
        except StopIteration:
            return [], [{'row': 0, 'message': 'The READINGS sheet is empty.'}]

        # Map header name -> column index (0-based), case/whitespace tolerant.
        header_index = {}
        for idx, cell in enumerate(headers):
            name = _text(cell)
            if name:
                header_index[name.upper()] = idx
        missing = [c for c in REQUIRED_COLUMNS if c not in header_index]
        if missing:
            return [], [{'row': 0,
                         'message': 'The sheet is missing the required column(s) '
                                    f'{", ".join(missing)} — download the template from this '
                                    'page and paste your data into it.'}]

        def cell(row_number, header):
            idx = header_index.get(header)
            return None if idx is None else row_values[idx]

        parsed_rows = []
        errors = []
        for row_number, row_values in enumerate(rows_iter, start=2):
            if not row_values or all(v is None or (isinstance(v, str) and not v.strip())
                                    for v in row_values):
                continue  # fully blank rows are skipped
            element = _text(cell(row_number, 'STRUCTURAL ELEMENT'))
            if element is None:
                errors.append({'row': row_number,
                               'message': 'STRUCTURAL ELEMENT is required — reference the '
                                          'target element for every row.'})
                continue
            test_type = _test_type(cell(row_number, 'TEST TYPE'), row_number, errors)
            if test_type is None:
                continue
            parsed_rows.append({
                'row': row_number,
                'element': element,
                'floor': _text(cell(row_number, 'FLOOR')) or '',
                'test_type': test_type,
                'point': _text(cell(row_number, 'POINT')),
                'path_length_mm': _number(cell(row_number, 'PATH LENGTH L (MM)'),
                                          row_number, 'PATH LENGTH L (MM)', errors),
                'transit_time_us': _number(cell(row_number, 'TRANSIT TIME T (US)'),
                                           row_number, 'TRANSIT TIME T (US)', errors),
                'uncracked_transit_time_us': _number(cell(row_number, 'T UNCRACKED (US)'),
                                                     row_number, 'T UNCRACKED (US)', errors),
                'surface_condition': _text(cell(row_number, 'SURFACE CONDITION')) or '',
                'transducer_frequency_khz': _number(
                    cell(row_number, 'TRANSDUCER FREQUENCY (KHZ)'),
                    row_number, 'TRANSDUCER FREQUENCY (KHZ)', errors),
                'transducer_type': _transducer_type(cell(row_number, 'TRANSDUCER TYPE')),
                'test_location': _text(cell(row_number, 'TEST LOCATION')) or '',
                'weather_condition': _text(cell(row_number, 'WEATHER CONDITION')) or '',
                'notes': _text(cell(row_number, 'NOTES')) or '',
            })
    finally:
        workbook.close()

    if errors:
        return [], errors
    if not parsed_rows:
        return [], [{'row': 0, 'message': 'No reading rows were found below the header — '
                                          'type your readings into the READINGS sheet (the '
                                          'EXAMPLE sheet is never imported).'}]

    # ---- per-type measurement completeness (row-level, before grouping) ----
    for row in parsed_rows:
        label = f"row {row['row']} ({row['element']})"
        if row['test_type'] == 'pulse_velocity':
            if row['path_length_mm'] is None:
                errors.append({'row': row['row'],
                               'message': f'{label}: PATH LENGTH L (MM) is required for a '
                                          'pulse-velocity point.'})
            if row['transit_time_us'] is None:
                errors.append({'row': row['row'],
                               'message': f'{label}: TRANSIT TIME T (US) is the field '
                                          'measurement — it cannot be blank.'})
        elif row['test_type'] == 'crack_depth':
            if row['path_length_mm'] is None:
                errors.append({'row': row['row'],
                               'message': f'{label}: PATH LENGTH L (MM) (the transducer '
                                          'spacing) is required for a crack-depth point.'})
            if row['transit_time_us'] is None or row['uncracked_transit_time_us'] is None:
                errors.append({'row': row['row'],
                               'message': f'{label}: crack-depth needs BOTH the cracked '
                                          '(TRANSIT TIME T) and uncracked (T UNCRACKED) '
                                          'transit times — they are field measurements.'})
        elif not row['surface_condition']:
            errors.append({'row': row['row'],
                           'message': f'{label}: SURFACE CONDITION is the field record for a '
                                      'surface-quality point — it cannot be blank.'})
    if errors:
        return [], errors

    # ---- group consecutive rows into tests ----
    groups = []
    seen_keys = {}
    for row in parsed_rows:
        key = (row['element'], row['test_type'], row['floor'])
        if groups and (groups[-1]['element'], groups[-1]['test_type'], groups[-1]['floor']) == key:
            groups[-1]['rows'].append(row)
            continue
        if key in seen_keys:
            first = seen_keys[key]
            errors.append({
                'rows': f"{row['row']}-{row['row']}",
                'message': f"Rows for element {row['element']!r} ({row['test_type']}, "
                           f"floor {row['floor'] or 'not recorded'}) appear in two separate "
                           f"blocks (first block started at row {first}). Keep one element's "
                           'points on consecutive rows so they form a single test.'})
            continue
        seen_keys[key] = row['row']
        groups.append({'element': row['element'], 'test_type': row['test_type'],
                       'floor': row['floor'], 'rows': [row]})
    if errors:
        return [], errors

    # ---- build serializer payloads ----
    for group in groups:
        first = group['rows'][0]
        readings = []
        for row in group['rows']:
            reading = {}
            if row['point']:
                reading['point_label'] = row['point']
            if row['test_type'] == 'pulse_velocity':
                reading['path_length_mm'] = row['path_length_mm']
                reading['transit_time_us'] = row['transit_time_us']
            elif row['test_type'] == 'crack_depth':
                reading['path_length_mm'] = row['path_length_mm']
                reading['transit_time_us'] = row['transit_time_us']
                reading['uncracked_transit_time_us'] = row['uncracked_transit_time_us']
            else:
                reading['surface_condition'] = row['surface_condition']
            if row['notes']:
                reading['notes'] = row['notes']
            readings.append(reading)

        payload = {
            'test_type': group['test_type'],
            'structural_element': group['element'],
            'floor': group['floor'],
            'transducer_type': first['transducer_type'],
            'test_location': first['test_location'],
            'weather_condition': first['weather_condition'],
            'readings': readings,
        }
        if first['transducer_frequency_khz'] is not None:
            payload['transducer_frequency_khz'] = int(first['transducer_frequency_khz'])
        if group['test_type'] == 'pulse_velocity':
            payload['path_length_mm'] = first['path_length_mm']
        elif group['test_type'] == 'crack_depth':
            payload['crack_path_length_mm'] = first['path_length_mm']
        group['payload'] = payload
        group['first_row'] = group['rows'][0]['row']
        group['last_row'] = group['rows'][-1]['row']
        del group['rows']
    return groups, []


def example_rows(sample_elements=None):
    """The EXAMPLE sheet's filled rows, in TEMPLATE_COLUMNS order.

    Exposed as the single definition of "the template's illustration" so the
    purge command can match registry records against the template's OWN
    example values instead of a copied list that could drift from it.

    ``sample_elements`` — optional ``(pulse_name, crack_name, surface_name)``
    of REAL elements from the project's imported BIM model, so the
    illustrations point at actual members; the fallback names are generic.
    """
    pulse_name, crack_name, surface_name = sample_elements or (
        'COL-A1', 'BEAM-B2', 'WALL-W1')
    return [
        # Format illustration only — this sheet is never imported. Pulse
        # velocity: consecutive rows -> ONE test with points A, B, C.
        [pulse_name, 'Ground Floor', 'Pulse Velocity', 'A', 120, 30.1, None, None,
         25, 'Direct', 'Grid B/4', 'Sunny',
         'EXAMPLE ONLY — copy the shape into READINGS, not the numbers'],
        [pulse_name, 'Ground Floor', 'Pulse Velocity', 'B', 120, 29.8, None, None,
         25, 'Direct', 'Grid B/4', 'Sunny', None],
        [pulse_name, 'Ground Floor', 'Pulse Velocity', 'C', 120, 30.3, None, None,
         25, 'Direct', 'Grid B/4', 'Sunny', None],
        # Crack depth: both the cracked AND uncracked transit times.
        [crack_name, 'First Floor', 'Crack Depth', 'A', 200, 52.1, 50.1, None,
         25, 'Direct', 'Grid D/2', 'Sunny', 'Hairline crack mid-span'],
        [crack_name, 'First Floor', 'Crack Depth', 'B', 200, 53.0, 50.2, None,
         25, 'Direct', 'Grid D/2', 'Sunny', None],
        # Surface quality: written condition instead of times.
        [surface_name, 'Ground Floor', 'Surface Quality', 'A', None, None, None,
         'Smooth, no honeycombing', None, None, 'East elevation', 'Sunny', None],
        [surface_name, 'Ground Floor', 'Surface Quality', 'B', None, None, None,
         'Minor voids near base', None, None, 'East elevation', 'Sunny', None],
    ]


def build_template_bytes(sample_elements=None):
    """The downloadable .xlsx template. The READINGS sheet ships EMPTY
    (headers only) — it is the only sheet the importer reads, so example
    data can never enter the registry. A separate EXAMPLE sheet shows
    filled rows for all three test types (clearly labelled as a format
    illustration, never imported), plus a HOW TO FILL sheet with the
    column instructions.

    ``sample_elements`` — optional ``(pulse_name, crack_name, surface_name)``
    of REAL elements from the project's imported BIM model. When given, the
    EXAMPLE rows reference those names so the illustrations point at actual
    members; the fallback names are generic."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'READINGS'
    header_font = Font(bold=True)
    for col, header in enumerate(TEMPLATE_COLUMNS, start=1):
        cell = sheet.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        sheet.column_dimensions[get_column_letter(col)].width = COLUMN_WIDTHS[header]
    sheet.freeze_panes = 'A2'

    example = workbook.create_sheet('EXAMPLE')
    for col, header in enumerate(TEMPLATE_COLUMNS, start=1):
        cell = example.cell(row=1, column=col, value=header)
        cell.font = header_font
        example.column_dimensions[get_column_letter(col)].width = COLUMN_WIDTHS[header]
    for row_number, row in enumerate(example_rows(sample_elements), start=2):
        for col, value in enumerate(row, start=1):
            if value is not None:
                example.cell(row=row_number, column=col, value=value)
    example.cell(
        row=len(example_rows(sample_elements)) + 3, column=1,
        value='EXAMPLE ONLY — this sheet is never imported. Only the READINGS '
              'sheet is read, so these rows can never reach the registry. '
              'Type your real readings into READINGS.')

    instructions = workbook.create_sheet('HOW TO FILL')
    lines = [
        ('NEXUCON — PUNDIT READINGS BULK UPLOAD', True),
        ('', False),
        ('Fill the READINGS sheet: ONE ROW PER TEST POINT. Upload it on the Data', False),
        ('Collection page (Batch Import tab) against the project — every value is', False),
        ('calculated by the platform; nothing is computed in the sheet.', False),
        ('', False),
        ('The EXAMPLE sheet shows filled rows for all three test types — it is a', False),
        ('format illustration ONLY and is NEVER imported. Copy its shape into', False),
        ('the READINGS sheet, never its numbers: only READINGS is read, so', False),
        ('example data cannot reach the registry.', False),
        ('', False),
        ('STRUCTURAL ELEMENT — the element being tested (reference the target element', False),
        ('   from the project\'s BIM model, e.g. Column C-102 or Floor:200THK RC SLAB).', False),
        ('   An element\'s consecutive rows form ONE test with points A, B, C...', False),
        ('FLOOR — e.g. Ground Floor, First Floor, Roof (optional but recommended —', False),
        ('   the report groups results by floor).', False),
        ('TEST TYPE — PULSE VELOCITY / CRACK DEPTH / SURFACE QUALITY.', False),
        ('POINT — point label (A, B, C...). Leave blank to auto-assign in row order.', False),
        ('PATH LENGTH L (MM) — the acoustic path / transducer spacing, a measurement.', False),
        ('TRANSIT TIME T (US) — the transit time at that point. For CRACK DEPTH this', False),
        ('   is the CRACKED path time.', False),
        ('T UNCRACKED (US) — CRACK DEPTH only: the uncracked-path transit time.', False),
        ('SURFACE CONDITION — SURFACE QUALITY only: the observed condition.', False),
        ('TRANSDUCER FREQUENCY (KHZ) / TRANSDUCER TYPE — optional; DIRECT is', False),
        ('   assumed when the type is blank.', False),
        ('TEST LOCATION / WEATHER CONDITION / NOTES — optional context.', False),
        ('', False),
        ('The import is all-or-nothing: if any row is rejected the file is rejected', False),
        ('with the row numbers and reasons, and nothing is written to the registry.', False),
        ('Velocity, quality grade, compressive strength and crack depth are always', False),
        ('computed by the platform — never typed into the sheet.', False),
    ]
    for row_number, (text, bold) in enumerate(lines, start=1):
        cell = instructions.cell(row=row_number, column=1, value=text)
        if bold:
            cell.font = Font(bold=True, size=13)
    instructions.column_dimensions['A'].width = 105
    return workbook


def build_template_response_bytes(sample_elements=None):
    """Serialised template workbook (.xlsx bytes) for the download endpoint."""
    import io
    buffer = io.BytesIO()
    build_template_bytes(sample_elements).save(buffer)
    return buffer.getvalue()
