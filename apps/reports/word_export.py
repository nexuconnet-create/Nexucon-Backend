"""
Word (.docx) export of the statutory NDT report (8 Sep 2026 review
meeting, item H7; 4 Sep register C5): an editable Word document carrying
the same sections, the same CMS-resolved prose and the same
server-computed figures as the fpdf2 PDF in ndt_reports.py.

The PDF remains the certified document of record (its exact bytes are
archived + checksummed); the .docx is the editable working copy the
client asked for. Every value here is read from the same helpers the
PDF uses (NDTReportService._element_data, ecs_report_disclosure,
report_cms.get_cms_text), so the two outputs cannot diverge.
"""
import io
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches

from .ndt_reports import (
    ECS_FORMULA_LINE, NDTReportService, _element_display,
    ecs_report_disclosure,
)
from .report_cms import cms_list_items, cms_paragraphs, get_cms_text


def _add_para(doc, text, *, bold=False, size=11, align=None, space_after=6):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = 'Cambria'
    if align is not None:
        para.alignment = align
    para.paragraph_format.space_after = Pt(space_after)
    return para


def _add_heading(doc, text, level=1):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = True
    run.font.size = Pt(14 if level == 1 else 12)
    run.font.name = 'Cambria'
    para.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    para.paragraph_format.space_after = Pt(6)
    return para


def _add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = ''
        run = cell.paragraphs[0].add_run(str(text))
        run.bold = True
        run.font.size = Pt(9)
        run.font.name = 'Cambria'
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = ''
            run = cell.paragraphs[0].add_run(str(value))
            run.font.size = Pt(9)
            run.font.name = 'Cambria'
    _add_para(doc, '', size=4, space_after=0)
    return table


class NDTWordExporter:
    """Builds the .docx edition of the statutory NDT report."""

    @classmethod
    def export_docx(cls, project, user=None):
        from apps.digital_eye.models import PUNDITTest, RebarTest

        S = NDTReportService
        tests = list(
            PUNDITTest.objects
            .filter(project=project)
            .select_related('device', 'operator')
            .prefetch_related('files', 'readings')
            .order_by('structural_element', 'tested_at')
        )
        rebar_tests = list(
            RebarTest.objects.filter(project=project).order_by('recorded_at')
        )
        pulse_tests = [t for t in tests if t.test_type == 'pulse_velocity']
        crack_tests = [t for t in tests if t.test_type == 'crack_depth']
        surface_tests = [t for t in tests if t.test_type == 'surface_quality']
        devices = []
        for t in tests:
            if t.device and t.device not in devices:
                devices.append(t.device)

        element_data = S._element_data(pulse_tests)
        good_members = [e for e in element_data if e['remark'] == 'GOOD']
        poor_members = [e for e in element_data if e['remark'] == 'POOR']
        visual_notes = S._visual_observations(tests)
        report_no, _year = S._effective_report_number(project, tests)
        tested = [t.test_date for t in tests if t.test_date]
        date_max = max(tested) if tested else datetime.now().date()

        # Generated-content CMS bodies (11 Sep): same resolution as the PDF
        # renderer, so the .docx can never show different wording.
        from apps.digital_eye.models import BIMElementMapping
        floors_present = sorted({e['floor_label'] for e in element_data})
        bim_levels = sorted(set(
            BIMElementMapping.objects.filter(project=project)
            .exclude(level='').values_list('level', flat=True)))
        has_drawings = BIMElementMapping.objects.filter(
            project=project).exists()
        same_day = bool(tested) and min(tested) == max(tested)
        date_min = min(tested) if tested else date_max
        cms = S._computed_bodies(
            project, tests=tests, rebar_tests=rebar_tests,
            element_data=element_data, good_members=good_members,
            poor_members=poor_members, visual_notes=visual_notes,
            floors_present=floors_present, bim_levels=bim_levels,
            has_drawings=has_drawings, tested=tested, same_day=same_day,
            date_min=date_min, date_max=date_max)

        doc = Document()
        # ------------------------------------------------------------ cover
        _add_para(doc, 'LAGOS STATE MATERIALS TESTING LABORATORY',
                  bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER,
                  space_after=18)
        _add_para(doc, 'IN-SITU INTEGRITY TEST (NON-DESTRUCTIVE) OF '
                       'COMPRESSIVE STRENGTH OF STRUCTURAL MEMBERS',
                  bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER,
                  space_after=18)
        _add_table(doc, ['FIELD', 'VALUE'], [
            ['Report No', report_no],
            ['Project', project.name or '-'],
            ['Client', project.client_name or '-'],
            ['Site Address', ', '.join(
                p for p in (project.site_address, project.lga,
                            project.state) if p) or '-'],
            ['Date of Test', date_max.strftime('%d/%m/%Y')],
        ])
        _add_para(doc, 'NOTE: this Word document is an editable working '
                       'copy. The certified document of record is the PDF '
                       'edition archived (SHA-256 checksummed) on the '
                       'platform.', size=9, space_after=12)

        # ------------------------------------------------------------ 1.0
        _add_heading(doc, '1.0 INTRODUCTION')
        for para in cms_paragraphs(get_cms_text(project, 'introduction')[0]):
            _add_para(doc, para)
        # Generated-content CMS section (11 Sep): the project paragraphs
        # resolve exactly as the PDF renders them.
        for para in cms_paragraphs(
                get_cms_text(project, 'introduction_project',
                             computed=cms)[0]):
            _add_para(doc, para)

        # ------------------------------------------------------------ 2.0
        _add_heading(doc, '2.0 PURPOSE OF INVESTIGATION')
        _add_para(doc, 'The purpose of the investigation is to:')
        for item in cms_list_items(get_cms_text(project,
                                                'purpose_items')[0]):
            para = doc.add_paragraph(item, style='List Number')
            para.paragraph_format.space_after = Pt(4)

        # ------------------------------------------------------------ 3.0
        _add_heading(doc, '3.0 LITERATURE REVIEW')
        for para in cms_paragraphs(get_cms_text(project,
                                                'literature_review')[0]):
            _add_para(doc, para)
        _add_para(doc, 'The pulse velocity is calculated using the '
                       'relationship:')
        _add_para(doc, 'UPV = L / t', align=WD_ALIGN_PARAGRAPH.CENTER)
        _add_para(doc, 'L = Distance between transducers (mm)', size=10)
        _add_para(doc, 't = Pulse transit time (microseconds)', size=10)
        _add_para(
            doc,
            'The testing procedure and interpretation of results are '
            'conducted in accordance with international standards including '
            'ASTM C597 (2020), BS EN 12504-4:2004 and ACI 228.2R (2018).')
        disclosure, derivation = ecs_report_disclosure(project)
        _add_para(doc, disclosure)
        _add_para(
            doc,
            'Derivation of the reported results: each test point velocity '
            'is V = L / t; the element pulse velocity is the arithmetic '
            'mean of its point velocities, V(element) = (V1 + V2 + ... + '
            f'Vn) / n; and the estimated compressive strength follows '
            f'{derivation}. Pulse velocities in the Section 5.0 tables are '
            'reported in metres per second (m/s); 1 km/s = 1000 m/s.')
        example = S._worked_example(project, element_data)
        if example:
            _add_para(doc, example)

        # ------------------------------------------------------------ 4.1
        _add_heading(doc, '4.1 VISUAL TEST')
        for para in cms_paragraphs(get_cms_text(project,
                                                'visual_preamble')[0]):
            _add_para(doc, para)
        # Generated-content CMS section (11 Sep): the recorded observations
        # resolve exactly as the PDF renders them.
        visual_body, _src = get_cms_text(project, 'visual_observations',
                                         computed=cms)
        if visual_body is not None:
            for item in cms_list_items(visual_body):
                _add_para(doc, item, size=10)
        else:
            _add_para(doc, 'No visual/surface condition observations '
                           'recorded.')

        # ------------------------------------------------------------ 4.2
        _add_heading(doc, '4.2 METHODOLOGY')
        for para in cms_paragraphs(
                get_cms_text(project, 'methodology_equipment',
                             computed=cms)[0]):
            _add_para(doc, para)
        _add_heading(doc, 'CONCRETE', level=2)
        for para in cms_paragraphs(get_cms_text(project,
                                                'methodology_concrete')[0]):
            _add_para(doc, para)
        _add_para(doc, 'The Pundit test equipment can also determine the '
                       'following:')
        for item in ('The homogeneity and uniformity of the concrete.',
                     'Changes in the strength of the concrete which may '
                     'occur with time.',
                     'The quality of the concrete in relation to standard '
                     'requirements.',
                     'The quality of one element of concrete in relation to '
                     'another.'):
            doc.add_paragraph(item, style='List Bullet')
        from apps.digital_eye.strength_curves import resolve_active_curve
        try:
            active_curve = resolve_active_curve(project)
        except Exception:  # noqa: BLE001 — conversion line must not kill export
            active_curve = None
        if active_curve is not None and active_curve.project_id:
            from apps.digital_eye.strength_curves import formula_display
            conversion = (
                formula_display(active_curve.curve_type,
                                active_curve.formula_params or {})
                + ' (f_cu in N/mm2, V in m/s'
                + (', R the rebound number)'
                   if active_curve.curve_type == 'sonreb' else ')'))
        else:
            conversion = ECS_FORMULA_LINE
        _add_para(doc, 'Estimated compressive strength conversion: '
                       + conversion + '. The full calibration statement is '
                       'given in Section 3.0.')

        # ------------------------------------------------------------ 4.3
        _add_heading(doc, '4.3 REINFORCING BAR (REBAR) ASSESSMENT')
        # Generated-content CMS section (11 Sep): resolves exactly as the
        # PDF renders it; without a survey the honest statement prints.
        rebar_body, _src = get_cms_text(project, 'rebar_statement',
                                        computed=cms)
        if rebar_body is not None:
            _add_para(doc, rebar_body)
            _add_table(
                doc,
                ['S/N', 'STRUCTURAL MEMBER', 'MAIN BAR (MM)', 'LINKS (MM)',
                 'SPACING (MM)', 'COVER DEPTH (MM)'],
                [[str(i + 1),
                  (rt.structural_element or 'Unknown').upper(),
                  str(int(rt.main_bar_mm)) if rt.main_bar_mm else '-',
                  str(int(rt.links_mm)) if rt.links_mm else '-',
                  str(rt.spacing_mm) if rt.spacing_mm else '-',
                  str(int(rt.cover_depth_mm)) if rt.cover_depth_mm else '-']
                 for i, rt in enumerate(rebar_tests)])
        else:
            _add_para(doc, 'No rebar assessment was recorded during this '
                           'investigation.')

        # ------------------------------------------------------------ 4.4
        _add_heading(doc, '4.4 EQUIPMENT')
        if devices:
            _add_table(doc, ['NAME OF EQUIPMENT', 'EQUIPMENT ID'],
                       [[d.name or d.model or '-', d.device_id or '-']
                        for d in devices])
        else:
            _add_para(doc, 'No field device was recorded against these '
                           'tests.')

        # ------------------------------------------------------------ 5.0
        _add_heading(doc, '5.0 ANALYSIS OF TEST RESULT')
        if crack_tests:
            _add_heading(doc, '5.1 CRACK DEPTH MEASUREMENTS '
                              '(TIME-DIFFERENCE METHOD)', level=2)
            for t in crack_tests:
                rows = t.reading_rows()
                mean_depth = S._crack_depth(t)
                _add_table(
                    doc,
                    ['ELEMENT', 'POINT', 'SPACING L (MM)', 'T CRACKED (US)',
                     'T UNCRACKED (US)', 'CRACK DEPTH (MM)',
                     'MEAN DEPTH (MM) / REMARK'],
                    [[_element_display(t.structural_element) if i == 0 else '',
                      r['label'] or '-',
                      S._fmt(r['path_mm'], 1),
                      S._fmt(r['transit_us'], 1),
                      S._fmt(r['uncracked_us'], 1),
                      S._fmt(r['crack_depth_mm'], 1),
                      (f"{S._fmt(mean_depth, 1)} / {S._crack_remark(t)}"
                       if i == len(rows) // 2 or len(rows) == 1 else '')]
                     for i, r in enumerate(rows)])
        if not element_data:
            _add_para(doc, 'No pulse velocity tests recorded for this '
                           'project.')
        else:
            _add_heading(doc, 'SUMMARY OF TEST ANALYSIS', level=2)
            analysis_groups = {}
            for e in element_data:
                key = (e['floor_label'], e['member_type'])
                g = analysis_groups.setdefault(key, {'count': 0, 'points': 0})
                g['count'] += 1
                g['points'] += e['n_points']
            _add_table(
                doc,
                ['STRUCTURAL MEMBER', 'NUMBER TESTED', 'LOCATION',
                 'NO OF POINT TAKEN'],
                [[member.title(), str(g['count']), floor.title(),
                  str(g['points'])]
                 for (floor, member), g in sorted(
                     analysis_groups.items())])

            _add_heading(doc, 'SUMMARY OF TEST RESULTS', level=2)
            for e in element_data:
                rows = e['rows']
                mid = len(rows) // 2 if len(rows) > 1 else 0
                remark = e['remark']
                if (e['spread_pct'] is not None and e['spread_pct'] > 2.0):
                    remark += (f" (POINT SPREAD "
                               f"{e['spread_km_s'] * 1000:.0f} M/S, "
                               f"±{e['spread_pct'] / 2:.1f}%)")
                _add_table(
                    doc,
                    ['Structural Element', 'PATH LENGTH', 'TRANSIT TIME',
                     'PULSE VELOCITY (M/S)', 'E.C.S',
                     'AVERAGE COMPRESSIVE STRENGTH (N/mm2)', 'REMARK'],
                    [[_element_display(e['element']) if i == 0 else '',
                      S._fp(r['path_mm']),
                      S._f1(r['transit_us']),
                      S._fms(r['velocity_km_s']),
                      S._f1(r['ecs_mpa']),
                      S._f1(e['mean_ecs']) if i == mid else '',
                      remark if i == mid else '']
                     for i, r in enumerate(rows)])
            result_groups = {}
            for e in element_data:
                key = (e['floor_label'], e['member_type'])
                g = result_groups.setdefault(
                    key, {'good': 0, 'poor': 0, 'total': 0})
                g['total'] += 1
                if e['remark'] == 'GOOD':
                    g['good'] += 1
                elif e['remark'] == 'POOR':
                    g['poor'] += 1
            _add_table(
                doc,
                ['STRUCTURAL MEMBER', 'LOCATION', 'GOOD (NO, %)',
                 'POOR (NO, %)'],
                [[member.title(), floor.title(),
                  f"{g['good']} ({round(g['good'] * 100 / g['total'], 1)}%)",
                  f"{g['poor']} ({round(g['poor'] * 100 / g['total'], 1)}%)"]
                 for (floor, member), g in sorted(result_groups.items())])
        if surface_tests:
            _add_heading(doc, '5.2 SURFACE QUALITY OBSERVATIONS', level=2)
            for t in surface_tests:
                for r in t.reading_rows():
                    condition = r.get('surface_condition') \
                        or 'condition not recorded'
                    _add_para(
                        doc,
                        f"{_element_display(t.structural_element)} "
                        f"{r['label'] or '-'}: {condition}")

        # ------------------------------------------------------------ 6.0
        _add_heading(doc, '6.0 RECOMMENDATION')
        if element_data:
            lead_in = get_cms_text(project, 'recommendation_preamble')[0]
            findings_body, _src = get_cms_text(project, 'findings_statement',
                                               computed=cms)
            _add_para(doc, lead_in.rstrip() + ' ' + findings_body)
        else:
            _add_para(doc, 'No pulse velocity results are available for '
                           'this project; no recommendation on concrete '
                           'quality can be made.')
        from apps.evidence.models import AIAnalysisRecord
        recommendations = []
        for record in (AIAnalysisRecord.objects
                       .filter(project=project, analysis_type='pundit')
                       .order_by('-created_at')[:20]):
            for rec in (record.recommendations or []):
                if isinstance(rec, dict):
                    line = f'[{str(rec.get("priority", "Routine")).upper()}] ' \
                           f'{rec.get("recommendation", "")}'
                else:
                    line = str(rec)
                if line not in recommendations:
                    recommendations.append(line)
        for rec in recommendations:
            doc.add_paragraph(rec, style='List Bullet')

        # ------------------------------------------------------------ 7.0
        _add_heading(doc, '7.0 CONCLUSION')
        if element_data:
            for para in cms_paragraphs(get_cms_text(project,
                                                    'conclusion_preamble')[0]):
                _add_para(doc, para)
            # Generated-content CMS section (11 Sep): the conclusion items
            # resolve exactly as the PDF renders them.
            for item in cms_list_items(
                    get_cms_text(project, 'conclusion_items',
                                 computed=cms)[0]):
                para = doc.add_paragraph(item, style='List Number')
                para.paragraph_format.space_after = Pt(4)
            _add_para(
                doc,
                'However, it is imperative to state clearly that '
                'non-adherence to the recommendation excludes the testing '
                'laboratory of any responsibility.')
            _add_para(
                doc,
                'The test assumed 25 N/mm2 as the strength of the '
                'structural members, however a substructure probe is '
                'required to ascertain the integrity of the building '
                'foundation.')
        else:
            _add_para(doc, 'No ultrasonic pulse velocity results are '
                           'available for this project; no conclusion on '
                           'concrete quality can be drawn.')

        # ---------------------------------------------------------- sign-off
        user_label = 'Unauthenticated'
        if user and getattr(user, 'is_authenticated', False):
            user_label = user.get_full_name() or user.email
        operators = []
        for t in tests:
            label = (t.operator_name
                     or ((t.operator.get_full_name() or t.operator.email)
                         if t.operator else None))
            if label and label not in operators:
                operators.append(label)
        tested_by = (operators[0] if operators else 'NOT RECORDED').upper()
        _add_table(doc, ['TESTED BY', 'APPROVED BY'],
                   [[tested_by, user_label.upper()]])

        # ------------------------------------------------------------- footer
        digest = S._statutory_digest(project, tests, report_no)
        _add_para(doc, 'REPORT INTEGRITY', bold=True)
        _add_para(
            doc,
            'Content digest (SHA-256 of the underlying test identifiers, '
            'results and reading rows):')
        _add_para(doc, digest, size=9)
        _add_para(
            doc,
            f'Generated {datetime.now().strftime("%d/%m/%Y %H:%M")} by '
            f'{user_label}.', size=9)

        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()
